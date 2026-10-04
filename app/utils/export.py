import io
import json
from fpdf import FPDF
from docx import Document
from docx.shared import Pt
from typing import Dict, Any

def _sanitise_for_pdf(text: str) -> str:
    """Replace characters that the default FPDF Latin-1 font cannot render."""
    replacements = {
        "\u2192": "->", "\u2190": "<-", "\u2194": "<->", "\u2022": "-", 
        "\u2026": "...", "\u201c": '"', "\u201d": '"', "\u2018": "'", 
        "\u2019": "'", "\u2013": "-", "\u2014": "--", "\u2260": "!=", 
        "\u2264": "<=", "\u2265": ">=", "\u00d7": "x", "\u00f7": "/", 
        "\u2713": "OK", "\u26a0": "(!)", "\u00a0": " ",
    }
    for char, replacement in replacements.items():
        text = text.replace(char, replacement)
    return text.encode("latin-1", errors="replace").decode("latin-1")

def generate_pdf(text: str, metadata: Dict[str, Any], structured_data: Dict[str, Any]) -> bytes:
    """Generate a clean PDF report."""
    pdf = FPDF()
    pdf.set_margins(left=20, top=20, right=20)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    # Title
    pdf.set_font("Helvetica", style="B", size=16)
    pdf.cell(0, 10, "Intelligent Document Recognition Report", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(8)

    # Metadata
    pdf.set_font("Helvetica", style="B", size=12)
    pdf.cell(0, 8, "Analysis Summary", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    pdf.cell(0, 6, f"Document Type: {metadata.get('doc_type', 'Unknown')}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, f"Model Used: {metadata.get('model', 'Unknown')}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, f"Processing Time: {metadata.get('time', '0')} seconds", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)

    # Structured Data
    if any(structured_data.values()):
        pdf.set_font("Helvetica", style="B", size=12)
        pdf.cell(0, 8, "Structured Information", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=10)
        for key, val in structured_data.items():
            if val:
                val_str = ", ".join(val) if isinstance(val, list) else str(val)
                pdf.cell(0, 6, f"- {key.title()}: {val_str}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)

    # Main Text
    pdf.set_font("Helvetica", style="B", size=12)
    pdf.cell(0, 8, "Recognized Text", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=11)
    
    safe_text = _sanitise_for_pdf(text)
    for line in safe_text.split("\n"):
        content = line if line.strip() else " "
        pdf.multi_cell(0, 6, content, new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())

def generate_docx(text: str, metadata: Dict[str, Any]) -> bytes:
    """Generate an editable Word document."""
    doc = Document()
    title_para = doc.add_heading("Extracted Document", level=1)
    if title_para.runs:
        title_para.runs[0].font.size = Pt(18)
        
    doc.add_paragraph(f"Document Type: {metadata.get('doc_type', 'Unknown')} | Model: {metadata.get('model', 'Unknown')}")
    doc.add_paragraph("")
    
    for line in text.split("\n"):
        p = doc.add_paragraph(line)
        if p.runs:
            p.runs[0].font.size = Pt(11)
            
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def generate_json(text: str, metadata: Dict[str, Any], structured_data: Dict[str, Any]) -> bytes:
    """Generate machine-readable JSON."""
    payload = {
        "document_type": metadata.get("doc_type"),
        "recognition_model": metadata.get("model"),
        "processing_time_seconds": metadata.get("time"),
        "recognition_confidence": metadata.get("confidence_label"),
        "text": text,
        "entities": structured_data
    }
    return json.dumps(payload, indent=2).encode('utf-8')
